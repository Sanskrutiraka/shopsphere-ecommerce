from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status as http_status
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Sum, Count
from datetime import timedelta

from accounts.permissions import IsAdminUser
from accounts.views import log_admin_action
from payments.models import Payment
from orders.models import Order


def get_date_range(request):
    from_date = request.query_params.get('start_date') or request.query_params.get('from')
    to_date = request.query_params.get('end_date') or request.query_params.get('to')
    if not from_date:
        from_date = (timezone.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    if not to_date:
        to_date = timezone.now().strftime('%Y-%m-%d')
    return from_date, to_date


class PaymentReportView(APIView):
    permission_classes = [IsAdminUser]

    def get(self, request):
        from_date, to_date = get_date_range(request)
        fmt = request.query_params.get('file_format') or request.query_params.get('format', 'json')
        export_format = 'excel' if fmt == 'excel' else 'pdf' if fmt == 'pdf' else None

        payments = Payment.objects.filter(
            created_at__date__gte=from_date,
            created_at__date__lte=to_date,
            status='SUCCESS'
        ).select_related('order', 'user').order_by('created_at')

        summary = payments.aggregate(
            total=Sum('amount'),
            count=Count('id')
        )

        data = [{
            'Transaction ID': p.transaction_id,
            'Order #': p.order.order_number,
            'Customer': p.user.get_full_name(),
            'Email': p.user.email,
            'Amount (₹)': float(p.amount),
            'Method': p.method,
            'Date': p.paid_at.strftime('%Y-%m-%d %H:%M') if p.paid_at else '',
        } for p in payments]

        if fmt == 'excel':
            response = self._excel_response(data, from_date, to_date, summary)
            log_admin_action(
                request.user,
                'REPORT_EXPORT',
                'PaymentReport',
                None,
                f'Exported payment report as EXCEL for {from_date} to {to_date} ({len(data)} records)',
                request,
            )
            return response
        elif fmt == 'pdf':
            response = self._pdf_response(data, from_date, to_date, summary)
            log_admin_action(
                request.user,
                'REPORT_EXPORT',
                'PaymentReport',
                None,
                f'Exported payment report as PDF for {from_date} to {to_date} ({len(data)} records)',
                request,
            )
            return response

        return Response({
            'from': from_date, 'to': to_date,
            'total_amount': float(summary['total'] or 0),
            'total_transactions': summary['count'],
            'payments': data
        })

    def _excel_response(self, data, from_date, to_date, summary, report_title='Payment Report', filename_prefix='payment_report'):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import io

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = report_title

        headers = list(data[0].keys()) if data else []
        total_columns = max(len(headers), 1)
        last_col_letter = get_column_letter(total_columns)

        title_fill = PatternFill(start_color='111827', end_color='111827', fill_type='solid')
        header_fill = PatternFill(start_color='F97316', end_color='F97316', fill_type='solid')
        odd_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        even_row_fill = PatternFill(start_color='FFF7ED', end_color='FFF7ED', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB'),
        )

        # Title
        ws.merge_cells(f'A1:{last_col_letter}1')
        ws['A1'] = f'{report_title} — {from_date} to {to_date}'
        ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
        ws['A1'].fill = title_fill
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.row_dimensions[1].height = 28

        ws['A2'] = f'Total: ₹{summary["total"] or 0} | Transactions: {summary["count"]}'
        ws['A2'].font = Font(italic=True, color='4B5563')
        ws.merge_cells(f'A2:{last_col_letter}2')
        ws['A2'].alignment = Alignment(horizontal='left')

        ws['A3'] = f'Generated at: {timezone.localtime().strftime("%Y-%m-%d %H:%M")}'
        ws['A3'].font = Font(size=10, color='6B7280')
        ws.merge_cells(f'A3:{last_col_letter}3')
        ws['A3'].alignment = Alignment(horizontal='left')

        # Headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        ws.row_dimensions[4].height = 22
        ws.freeze_panes = 'A5'

        # Data
        for row_num, row in enumerate(data, 5):
            fill = even_row_fill if row_num % 2 == 0 else odd_row_fill
            for col_num, value in enumerate(row.values(), 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.fill = fill
                cell.border = thin_border
                header = headers[col_num - 1] if headers else ''
                if '₹' in header or 'Total' in header or 'Amount' in header:
                    try:
                        cell.value = float(value)
                        cell.number_format = '"₹"#,##0.00'
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    except (TypeError, ValueError):
                        cell.alignment = Alignment(horizontal='left', vertical='center')
                elif 'Date' in header:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        if headers:
            ws.auto_filter.ref = f'A4:{last_col_letter}{max(4, len(data) + 4)}'

        # Auto-width for data headers only (skip merged title row cells).
        for index, header in enumerate(headers, 1):
            max_len = len(str(header or ''))
            for row in data:
                max_len = max(max_len, len(str(row.get(header, '') or '')))
            ws.column_dimensions[get_column_letter(index)].width = min(max_len + 4, 40)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}_{from_date}_{to_date}.xlsx"'
        return response

    def _pdf_response(self, data, from_date, to_date, summary, report_title='Payment Report', filename_prefix='payment_report'):
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.units import inch
        import io

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=0.5*inch, bottomMargin=0.5*inch)
        styles = getSampleStyleSheet()
        orange_color = colors.HexColor('#F97316')
        elements = []

        title_style = ParagraphStyle('Title', parent=styles['Heading1'], textColor=orange_color, fontSize=16)
        elements.append(Paragraph(f'ShopSphere — {report_title}', title_style))
        subtitle_style = ParagraphStyle('SubTitle', parent=styles['Normal'], textColor=colors.HexColor('#4B5563'))
        elements.append(Paragraph(f'Period: {from_date} to {to_date}', subtitle_style))
        elements.append(Paragraph(f'Total: INR {summary["total"] or 0} | Transactions: {summary["count"]}', subtitle_style))
        elements.append(Spacer(1, 0.2*inch))

        summary_table = Table([
            ['Metric', 'Value'],
            ['Total Amount', f'INR {summary["total"] or 0}'],
            ['Total Records', f'{summary["count"]}'],
        ], colWidths=[2.2 * inch, 2.2 * inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#111827')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF7ED')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 0.25*inch))

        if data:
            headers = [str(h).replace('₹', 'INR') for h in data[0].keys()]
            table_data = [headers] + [[str(row.get(h, '')).replace('₹', 'INR') for h in data[0].keys()] for row in data]
            available_width = landscape(A4)[0] - (doc.leftMargin + doc.rightMargin)
            char_weights = []
            for header in headers:
                max_len = len(str(header))
                for row in data:
                    max_len = max(max_len, len(str(row.get(header, ''))))
                char_weights.append(min(max_len, 35) + 2)

            total_weight = sum(char_weights) or 1
            col_widths = [(available_width * (w / total_weight)) for w in char_weights]

            table = Table(table_data, repeatRows=1, colWidths=col_widths)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), orange_color),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF7ED')]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
                ('FONTSIZE', (0, 1), (-1, -1), 8.5),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ]))

            for index, header in enumerate(headers):
                if '₹' in header or 'Total' in header or 'Amount' in header:
                    table.setStyle(TableStyle([('ALIGN', (index, 1), (index, -1), 'RIGHT')]))
                if 'Date' in header:
                    table.setStyle(TableStyle([('ALIGN', (index, 1), (index, -1), 'CENTER')]))

            elements.append(table)
        else:
            elements.append(Paragraph('No data available for the selected period.', styles['Normal']))

        doc.build(elements)
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}_{from_date}_{to_date}.pdf"'
        return response


class OrderReportView(APIView):
    permission_classes = [IsAdminUser]

    def get(self, request):
        from_date, to_date = get_date_range(request)
        fmt = request.query_params.get('file_format') or request.query_params.get('format', 'json')

        orders = Order.objects.filter(
            placed_at__date__gte=from_date,
            placed_at__date__lte=to_date
        ).select_related('user').order_by('placed_at')

        summary = orders.aggregate(
            total_revenue=Sum('total_amount'),
            total_orders=Count('id')
        )
        status_breakdown = orders.values('status').annotate(count=Count('id'))

        data = [{
            'Order #': o.order_number,
            'Customer': o.user.get_full_name(),
            'Email': o.user.email,
            'Status': o.status,
            'Payment Method': o.payment_method,
            'Total (₹)': float(o.total_amount),
            'Date': o.placed_at.strftime('%Y-%m-%d %H:%M'),
        } for o in orders]

        if fmt == 'excel':
            view = PaymentReportView()
            response = view._excel_response(data, from_date, to_date, {
                'total': summary['total_revenue'], 'count': summary['total_orders']
            }, report_title='Order Report', filename_prefix='order_report')
            log_admin_action(
                request.user,
                'REPORT_EXPORT',
                'OrderReport',
                None,
                f'Exported order report as EXCEL for {from_date} to {to_date} ({len(data)} records)',
                request,
            )
            return response
        elif fmt == 'pdf':
            view = PaymentReportView()
            response = view._pdf_response(data, from_date, to_date, {
                'total': summary['total_revenue'], 'count': summary['total_orders']
            }, report_title='Order Report', filename_prefix='order_report')
            log_admin_action(
                request.user,
                'REPORT_EXPORT',
                'OrderReport',
                None,
                f'Exported order report as PDF for {from_date} to {to_date} ({len(data)} records)',
                request,
            )
            return response

        return Response({
            'from': from_date, 'to': to_date,
            'total_orders': summary['total_orders'],
            'total_revenue': float(summary['total_revenue'] or 0),
            'status_breakdown': list(status_breakdown),
            'orders': data,
        })
